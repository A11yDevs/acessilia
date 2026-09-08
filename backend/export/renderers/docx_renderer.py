from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from backend.i18n import t
from backend.pipeline.table_ast import MSG_TABLE_TEXT_CAPTION
from backend.pipeline.table_ast import rows_from_table_ast
from backend.pipeline.table_ast import table_ast_from_block
from backend.pipeline.verbosity_manager import filter_blocks_for_profile


def render_docx(
    document: dict[str, Any],
    output_path: Path,
    profile_name: str = "docx",
    filename: str = "",
) -> Path:
    """Renders the canonical document into a Word .docx file with styled headings, code and tables.

    Args:
        document (dict): Canonical document mapping (sections, optional title) to render.
        output_path (Path): Destination file path; the parent directory is created when missing.
        profile_name (str): Export verbosity profile applied to block filtering (default "docx").
        filename (str): Optional document heading text rendered as a centered level-1 heading (default "").

    Returns:
        Path: The output_path where the .docx file was written.
    """
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)
    if filename:
        heading = doc.add_heading(filename, level=1)
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for section in document.get("sections", []):
        _render_section(doc, section, profile_name)
    doc.save(str(output_path))
    return output_path


def _render_section(doc: Document, section: dict[str, Any], profile_name: str) -> None:
    """Appends a section heading, its profile-filtered blocks, and its nested children to the docx.

    Args:
        doc (Document): The python-docx Document being built (mutated in place).
        section (dict): Canonical section mapping with optional "title", "level", "blocks", "children".
        profile_name (str): Export verbosity profile applied to the section's block filtering.
    """
    if section.get("title"):
        doc.add_heading(section["title"], level=min(section.get("level", 1), 9))
    for block in filter_blocks_for_profile(section.get("blocks", []), profile_name):
        _render_block(doc, block)
    for child in section.get("children", []):
        _render_section(doc, child, profile_name)


def _render_block(doc: Document, block: dict[str, Any]) -> None:
    """Appends the docx content for a single canonical block, choosing markup by block type.

    Args:
        doc (Document): The python-docx Document being built (mutated in place).
        block (dict): Canonical block mapping with at least "type" and the type-specific payload.
    """
    block_type = block.get("type")
    if block_type == "heading":
        doc.add_heading(
            block.get("title", block.get("text", "")),
            level=min(block.get("level", 1), 9),
        )
    elif block_type == "paragraph":
        doc.add_paragraph(block.get("text", ""))
    elif block_type == "code":
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(block.get("text", ""))
        run.font.name = "Courier New"
        run.font.size = Pt(10)
        paragraph.style = doc.styles["No Spacing"]
    elif block_type == "list":
        style = "List Number" if block.get("ordered") else "List Bullet"
        for item in block.get("items", []):
            doc.add_paragraph(str(item), style=style)
    elif block_type == "table":
        table_ast = table_ast_from_block(block)
        rows = rows_from_table_ast(table_ast) if table_ast else block.get("rows", [])
        if rows:
            caption = table_ast.get("caption") if isinstance(table_ast, dict) else None
            if isinstance(caption, str) and caption.strip():
                doc.add_paragraph(t(MSG_TABLE_TEXT_CAPTION).format(caption=caption.strip()))
            table = doc.add_table(rows=len(rows), cols=max(len(row) for row in rows))
            table.style = "Table Grid"
            for i, row in enumerate(rows):
                for j, cell in enumerate(row):
                    table.cell(i, j).text = str(cell)
    elif block_type in {"details", "note", "warning", "quote", "image", "math"}:
        doc.add_paragraph(block.get("text", block.get("alt_text", "")))
    else:
        doc.add_paragraph(block.get("text", ""))
